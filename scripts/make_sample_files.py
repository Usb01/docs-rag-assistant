"""Builds the PDF and DOCX sample documents so the demo covers every supported format."""
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).resolve().parent.parent / "sample_docs"

SLA = [
    ("Falcon Freight – Business Service Level Agreement", None),
    ("1. Scope", "This agreement covers business account customers using Standard, Priority and Enterprise "
                 "support plans. It describes response times, delivery commitments and service credits."),
    ("2. Support plans",
     "Standard plan: support by email, first response within 8 business hours. "
     "Priority plan: email and phone support, first response within 2 business hours, named account manager. "
     "Enterprise plan: 24/7 phone line, first response within 30 minutes for critical issues, quarterly business reviews."),
    ("3. On-time delivery commitment",
     "Falcon Freight commits to delivering at least 95% of Standard domestic shipments within the published "
     "delivery window each calendar month. Enterprise customers get a 97% commitment."),
    ("4. Service credits",
     "If on-time delivery falls below the commitment in a month, the customer receives a credit on the next invoice: "
     "5% of that month's freight charges for each full percentage point below the commitment, capped at 25%. "
     "Credits must be requested within 30 days of the end of the month. Credits are the customer's sole remedy for late delivery."),
    ("5. Exclusions",
     "Delays caused by customs holds, incorrect addresses provided by the customer, force majeure events, "
     "public holidays, or shipments containing prohibited items do not count towards the on-time calculation."),
    ("6. Incident severity",
     "Critical: tracking or booking systems unavailable for all users. High: a single customer cannot book or "
     "track shipments. Normal: questions, reporting requests and minor issues. Critical incidents are updated "
     "every hour until resolved."),
    ("7. Term and termination",
     "The agreement runs for 12 months and renews automatically. Either party can terminate with 60 days' "
     "written notice. Outstanding invoices remain payable after termination."),
]

SECURITY = [
    ("IT Security Policy", "h"),
    ("Passwords and accounts", "h2"),
    ("Passwords must be at least 12 characters long. Multi-factor authentication is mandatory for email, the HR "
     "portal and the booking system. Never share your account, including with your manager.", "p"),
    ("Laptops and phones", "h2"),
    ("Company laptops are encrypted and must be locked when unattended. Personal phones may be used for email "
     "and Teams only after enrolling them in the company device management app. Lost or stolen devices must be "
     "reported to the IT helpdesk within 1 hour so they can be wiped remotely.", "p"),
    ("Reporting incidents", "h2"),
    ("Report suspicious emails with the 'Report phishing' button in Outlook. Any suspected security incident must "
     "be reported to security@falconfreight-demo.com or extension 4444 immediately. Do not try to investigate yourself.", "p"),
    ("Data handling", "h2"),
    ("Customer data, including names, addresses and phone numbers, may only be stored in approved systems. "
     "Exporting customer lists to personal email or USB drives is prohibited. Files shared outside the company "
     "must use the secure share link, which expires after 7 days.", "p"),
    ("Software", "h2"),
    ("Only software from the company portal may be installed. Requests for new software go through the IT helpdesk "
     "and are reviewed within 5 business days.", "p"),
]


def build_pdf():
    styles = getSampleStyleSheet()
    story = []
    for title, body in SLA:
        story.append(Paragraph(title, styles["Title" if body is None else "Heading2"]))
        if body:
            story.append(Paragraph(body, styles["BodyText"]))
            story.append(Spacer(1, 8))
    SimpleDocTemplate(str(OUT / "service_level_agreement.pdf"), pagesize=A4,
                      title="Falcon Freight SLA").build(story)


def build_docx():
    doc = Document()
    for text, kind in SECURITY:
        if kind == "h":
            doc.add_heading(text, 0)
        elif kind == "h2":
            doc.add_heading(text, 1)
        else:
            doc.add_paragraph(text)
    doc.save(OUT / "it_security_policy.docx")


if __name__ == "__main__":
    build_pdf()
    build_docx()
    print("written to", OUT)
